"""Opt-in integration tests against an isolated DB on a real MySQL server."""
import os, sys, uuid, secrets, subprocess, socket, time, io
from pathlib import Path
from datetime import datetime, timezone, timedelta
from concurrent.futures import ThreadPoolExecutor
import bcrypt, jwt, requests, pytest
from mysql_test_support import IsolatedDatabase
from PIL import Image

pytestmark=pytest.mark.skipif(not os.environ.get('MYSQL_TEST_URL'),reason='Set MYSQL_TEST_URL to an isolated MySQL test account; never uses a public demo')


@pytest.fixture(scope='module')
def live(tmp_path_factory):
    root=Path(__file__).resolve().parents[2]
    temporary=tmp_path_factory.mktemp('real-transactions')
    name='audit_mysql_'+uuid.uuid4().hex
    db=IsolatedDatabase(os.environ['MYSQL_TEST_URL'],name)
    company='company-'+uuid.uuid4().hex
    password=secrets.token_urlsafe(24)
    users={}
    for label,role in [('admin','admin'),('owner','owner'),('officer','petugas'),('other','petugas'),('outsider','petugas'),('foreign','admin')]:
        person={'id':uuid.uuid4().hex,'company_id':company if label!='foreign' else 'foreign-company','role':role,'name':'QA '+label,'email':label+'.'+name+'@example.com','status':'aktif','petugas_code':label,'password_hash':bcrypt.hashpw(password.encode(),bcrypt.gensalt()).decode()}
        db.users.insert_one(person);users[label]=person
    import base64
    signature=io.BytesIO();Image.new('RGB',(32,16),'blue').save(signature,format='PNG')
    db.companies.insert_one({'id':company,'nama':'Production QA','company_code':'QA','app_name':'BIG Mobile','director_signature':'data:image/png;base64,'+base64.b64encode(signature.getvalue()).decode(),'director_name':'QA Director'})
    secret=secrets.token_urlsafe(48)
    env={**os.environ,'APP_ENV':'production','MYSQL_URL':db.uri,'JWT_SECRET':secret,'DB_TRANSACTIONS':'true','FILE_STORAGE_BACKEND':'local','LOCAL_STORAGE_DIR':str(temporary/'uploads'),'APP_BASE_URL':'https://qa.example.com','CORS_ORIGINS':'https://qa.example.com','SEED_DEMO_DATA':'false','MONITOR_TOKEN':secrets.token_urlsafe(32)}
    with socket.socket() as sock:sock.bind(('127.0.0.1',0));web_port=sock.getsockname()[1]
    env['CORS_ORIGINS']+=f',https://localhost:{web_port}'
    env['ALLOW_INSECURE_LOCAL_TEST_DATABASE']='true'
    with socket.socket() as sock:sock.bind(('127.0.0.1',0));port=sock.getsockname()[1]
    output=(temporary/'server.log').open('w')
    process=subprocess.Popen([sys.executable,'-m','uvicorn','server:app','--host','127.0.0.1','--port',str(port),'--no-access-log'],cwd=root/'backend',env=env,stdout=output,stderr=output)
    url=f'http://127.0.0.1:{port}/api'
    tokens={label:jwt.encode({'sub':user['id'],'role':user['role'],'type':'access','ver':0,'iss':'big-mobile-api','aud':'big-mobile','jti':uuid.uuid4().hex,'iat':datetime.now(timezone.utc),'exp':datetime.now(timezone.utc)+timedelta(hours=1)},secret,algorithm='HS256') for label,user in users.items()}
    def call(who,method,path,expected=200,**kwargs):
        headers={**kwargs.pop('headers',{}),'Authorization':'Bearer '+tokens[who]}
        response=requests.request(method,url+path,headers=headers,timeout=30,**kwargs)
        assert response.status_code==expected,f'{method} {path}: {response.status_code} {response.text[:200]}'
        return response
    try:
        for _ in range(int(os.environ.get('TEST_DATABASE_READY_SECONDS', '60'))):
            if process.poll() is not None:raise RuntimeError('Integration server failed; inspect isolated server.log')
            try:
                if requests.get(url+'/health',timeout=2).status_code==200:break
            except requests.RequestException:pass
            time.sleep(1)
        else:raise RuntimeError('Integration readiness timed out')
        yield {'db':db,'users':users,'call':call,'url':url,'tokens':tokens,'env':env,'temporary':temporary,'root':root,'password':password,'web_port':web_port}
    finally:
        process.terminate()
        try:process.wait(timeout=15)
        except subprocess.TimeoutExpired:process.kill();process.wait()
        output.close()
        db.close()


def unit(live):
    call=live['call'];suffix=uuid.uuid4().hex
    client=call('admin','POST','/clients',json={'nama_perusahaan':'QA '+suffix}).json()
    sk=call('admin','POST','/surat-kuasa',json={'nomor':suffix,'client_id':client['id'],'tanggal_berlaku':'2020-01-01','tanggal_berakhir':'2099-12-31'}).json()
    account=call('admin','POST','/akun',json={'nomor_kontrak':suffix,'nama_debitur':'QA Debitur','client_id':client['id'],'surat_kuasa_id':sk['id']}).json()
    return account,sk


def assign(live,account):
    letter=live['call']('admin','POST','/penugasan',json={'account_id':account['id'],'petugas_id':live['users']['officer']['id']}).json()
    return live['call']('owner','POST','/penugasan/'+letter['assignment_id']+'/approve').json()


def report_payload(letter,account):
    return {'client_report_id':str(uuid.uuid4()),'payment_amount':'500000','assignment_id':letter['assignment_id'],'account_id':account['id'],'status':'TIDAK_DITEMUKAN','catatan':'Kunjungan sudah dilaksanakan dengan lengkap','latitude':'-6.2','longitude':'106.8'}


def report_files():
    output=io.BytesIO();Image.new('RGB',(32,32),'blue').save(output,'PNG')
    return {'photos': ('proof.png',output.getvalue(),'image/png')}


def test_assignment_failure_rolls_back_letter_unit_counter_and_audit(live):
    account,_=unit(live);db=live['db']
    audits=db.audit_logs.count_documents({})
    counters=list(db.counters.find({}))
    db.block_writes('assignment_letters',True)
    try:
        live['call']('admin','POST','/penugasan',500,json={'account_id':account['id'],'petugas_id':live['users']['officer']['id']})
        assert db.assignments.count_documents({'account_id':account['id']})==0
        assert db.accounts.find_one({'id':account['id']})['status']=='BELUM_DITUGASKAN'
        assert db.audit_logs.count_documents({})==audits
        assert list(db.counters.find({}))==counters
    finally:db.block_writes('assignment_letters',False)


def test_concurrent_assignment_only_one_commits(live):
    account,_=unit(live)
    body={'account_id':account['id'],'petugas_id':live['users']['officer']['id']}
    def send():return requests.post(live['url']+'/penugasan',json=body,headers={'Authorization':'Bearer '+live['tokens']['admin']},timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:send(),range(2)))
    assert sorted(response.status_code for response in responses)==[200,400]
    assert live['db'].assignments.count_documents({'account_id':account['id'],'status':'MENUNGGU_APPROVAL'})==1


def test_idempotent_assignment_returns_original_result(live):
    account,_=unit(live)
    body={'account_id':account['id'],'petugas_id':live['users']['officer']['id']}
    headers={'Idempotency-Key':uuid.uuid4().hex}
    first=live['call']('admin','POST','/penugasan',json=body,headers=headers).json()
    second=live['call']('admin','POST','/penugasan',json=body,headers=headers).json()
    assert first['assignment_id']==second['assignment_id']
    assert live['db'].assignments.count_documents({'account_id':account['id']})==1


def test_real_photo_report_review_isolation_and_renewal(live):
    account,sk=unit(live);letter=assign(live,account)
    image=io.BytesIO();Image.new('RGB',(32,32),'blue').save(image,format='PNG')
    payload=report_payload(letter,account);payload['status']='UNIT_DITEMUKAN'
    report=live['call']('admin','POST','/laporan',data=payload,files={'photos':('proof.png',image.getvalue(),'image/png')}).json()
    path=report['photos'][0]['url'][4:]
    assert live['call']('admin','GET',path).content==image.getvalue()
    live['call']('other','GET',path,404);live['call']('foreign','GET',path,404)
    live['call']('foreign','GET','/laporan/'+report['id'],404)
    live['call']('admin','PATCH','/laporan/'+report['id']+'/review',json={'catatan_admin':'Bukti sudah direview'})
    history=live['call']('officer','GET','/my/riwayat').json()
    assert next(item for item in history if item['id']==report['id'])['catatan_admin']=='Bukti sudah direview'
    live['call']('admin','PATCH','/surat-tugas/'+letter['id']+'/status',json={'status':'selesai'})
    new_sk=live['call']('admin','POST','/surat-kuasa',json={'nomor':uuid.uuid4().hex,'client_id':account['client_id']}).json()
    updated={key:value for key,value in account.items() if key in {'nomor_kontrak','nama_debitur','client_id','surat_kuasa_id'}}
    updated.update({'surat_kuasa_id':new_sk['id'],'nama_debitur':'Updated current master'})
    live['call']('admin','PUT','/akun/'+account['id'],json=updated)
    assert live['call']('admin','GET','/laporan/'+report['id']).json()['account']['nama_debitur']=='QA Debitur'
    assert live['db'].upload_intents.count_documents({'company_id':account['company_id']})==0


def test_report_metadata_failure_rolls_back_and_orphan_file_is_cleaned(live):
    account,_=unit(live);letter=assign(live,account);db=live['db']
    image=io.BytesIO();Image.new('RGB',(32,32),'blue').save(image,format='PNG')
    db.block_writes('report_photos',True)
    try:
        payload=report_payload(letter,account);payload['status']='UNIT_DITEMUKAN'
        live['call']('admin','POST','/laporan',500,data=payload,files={'photos':('proof.png',image.getvalue(),'image/png')})
        assert db.field_reports.count_documents({'assignment_id':letter['assignment_id']})==0
        assert db.accounts.find_one({'id':account['id']})['status']=='DITUGASKAN'
        intents=list(db.upload_intents.find({}))
        assert intents
        db.upload_intents.update_many({}, {'$set':{'created_at':datetime.now(timezone.utc)-timedelta(hours=2)}})
        code='import asyncio, server; asyncio.run(server.reconcile_uploads())'
        subprocess.run([sys.executable,'-c',code],cwd=live['root']/'backend',env=live['env'],check=True,capture_output=True,timeout=30)
        assert db.upload_intents.count_documents({})==0
        assert not (Path(live['env']['LOCAL_STORAGE_DIR'])/intents[0]['storage_path']).exists()
    finally:db.block_writes('report_photos',False)


def test_concurrent_reports_and_cancellation_remain_consistent(live):
    account,_=unit(live);letter=assign(live,account);payload=report_payload(letter,account)
    def send():return requests.post(live['url']+'/laporan',data=payload,files=report_files(),headers={'Authorization':'Bearer '+live['tokens']['admin']},timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:send(),range(2)))
    assert sorted(response.status_code for response in responses)==[200,200]
    assert live['db'].field_reports.count_documents({'assignment_id':letter['assignment_id']})==1
    account2,_=unit(live);letter2=assign(live,account2)
    def cancellation():return live['call']('admin','PATCH','/surat-tugas/'+letter2['id']+'/status',json={'status':'dibatalkan'})
    def reporting():return requests.post(live['url']+'/laporan',data=report_payload(letter2,account2),files=report_files(),headers={'Authorization':'Bearer '+live['tokens']['admin']},timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:
        close=pool.submit(cancellation);report=pool.submit(reporting);close.result();outcome=report.result()
    assert outcome.status_code in {200,400,409}
    assert live['db'].assignments.find_one({'id':letter2['assignment_id']})['status']=='DIBATALKAN'
    assert live['db'].assignments.count_documents({'account_id':account2['id'],'active_key':{'$exists':True}})==0


def test_logout_revokes_access_and_production_login_sets_secure_cookie(live):
    response=requests.post(live['url']+'/auth/login',json={'email':live['users']['other']['email'],'password':live['password']},timeout=30)
    assert response.status_code==200 and 'token' not in response.json()
    assert 'HttpOnly' in response.headers['set-cookie'] and 'Secure' in response.headers['set-cookie']
    live['call']('other','POST','/auth/logout')
    live['call']('other','GET','/auth/me',401)
    # Later ownership tests need an independently authenticated session.
    payload=jwt.decode(live['tokens']['other'],live['env']['JWT_SECRET'],algorithms=['HS256'],audience='big-mobile',issuer='big-mobile-api')
    payload['jti']=uuid.uuid4().hex
    live['tokens']['other']=jwt.encode(payload,live['env']['JWT_SECRET'],algorithm='HS256')


def test_paginated_master_data_and_historical_report_search(live):
    company=live['users']['admin']['company_id']
    live['db'].clients.insert_many([{'id':uuid.uuid4().hex,'company_id':company,'nama_perusahaan':'Pagination QA','status':'aktif','created_at':'2099-01-01'} for _ in range(3)])
    first=live['call']('admin','GET','/clients?search=Pagination%20QA&limit=2')
    second=live['call']('admin','GET','/clients?search=Pagination%20QA&limit=2&page=2')
    assert len(first.json())==2 and len(second.json())==1
    assert first.headers['X-Total-Count']=='3' and first.headers['X-Next-Page']=='2'
    assert len({item['id'] for item in first.json()+second.json()})==3
    reports=live['call']('admin','GET','/laporan?search=QA%20Debitur&limit=1')
    assert len(reports.json())==1
    assert reports.json()[0]['nama_debitur']=='QA Debitur'


def test_real_pdf_access_and_company_upload_audit_rollback(live):
    from pypdf import PdfWriter
    account,sk=unit(live)
    stream=io.BytesIO();writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.write(stream)
    uploaded=live['call']('admin','POST','/surat-kuasa/'+sk['id']+'/file',files={'file':('authority.pdf',stream.getvalue(),'application/pdf')}).json()
    path=uploaded['file_url'][4:]
    assert live['call']('admin','GET',path).content==stream.getvalue()
    live['call']('officer','GET',path,404)
    assign(live,account)
    assert live['call']('officer','GET','/my/tugas/'+account['id']).json()['surat_kuasa_file']==uploaded['file_url']
    assert live['call']('officer','GET',path).content==stream.getvalue()
    live['call']('foreign','GET',path,404)
    image=io.BytesIO();Image.new('RGBA',(32,32),(0,120,255,128)).save(image,format='PNG')
    db=live['db'];before=db.companies.find_one({'id':live['users']['admin']['company_id']}).get('logo')
    db.block_writes('audit_logs',True)
    try:
        live['call']('admin','POST','/company/logo',500,files={'file':('logo.png',image.getvalue(),'image/png')})
        assert db.companies.find_one({'id':live['users']['admin']['company_id']}).get('logo')==before
    finally:db.block_writes('audit_logs',False)
    uploaded=live['call']('admin','POST','/company/logo',files={'file':('logo.png',image.getvalue(),'image/png')}).json()
    assert uploaded['logo'].startswith('data:image/png;base64,')


@pytest.mark.skipif(os.environ.get('RUN_BROWSER_TESTS')!='true',reason='Set RUN_BROWSER_TESTS=true after building frontend and installing Playwright browser')
def test_https_browser_workflow_and_network_recovery(live):
    import json, ipaddress
    from cryptography import x509
    from cryptography.x509.oid import NameOID
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
    name=x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,'localhost')]);now=datetime.now(timezone.utc)
    certificate=x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key()).serial_number(x509.random_serial_number()).not_valid_before(now-timedelta(minutes=1)).not_valid_after(now+timedelta(days=1)).add_extension(x509.SubjectAlternativeName([x509.DNSName('localhost'),x509.IPAddress(ipaddress.ip_address('127.0.0.1'))]),critical=False).sign(key,hashes.SHA256())
    temporary=live['temporary'];key_path=temporary/'qa.key';cert_path=temporary/'qa.crt';image=temporary/'photo.png'
    from pypdf import PdfWriter
    pdf=temporary/'authority.pdf';writer=PdfWriter();writer.add_blank_page(width=100,height=100)
    with pdf.open('wb') as output:writer.write(output)
    key_path.write_bytes(key.private_bytes(serialization.Encoding.PEM,serialization.PrivateFormat.PKCS8,serialization.NoEncryption()));cert_path.write_bytes(certificate.public_bytes(serialization.Encoding.PEM));Image.new('RGB',(64,64),'blue').save(image)
    fixture=temporary/'browser-fixture.json'
    print_pdf=temporary/'printed-letter.pdf'
    fixture.write_text(json.dumps({'api_url':live['url'],'web_port':live['web_port'],'key':str(key_path),'cert':str(cert_path),'image':str(image),'pdf':str(pdf),'print_pdf':str(print_pdf),'password':live['password'],'admin_email':live['users']['admin']['email'],'officer_email':live['users']['officer']['email'],'teammate_email':live['users']['other']['email'],'admin_token':live['tokens']['admin'],'owner_token':live['tokens']['owner'],'owner_email':live['users']['owner']['email']}))
    try:
        result=subprocess.run(['node',str(live['root']/'frontend/e2e/workflow.mjs')],env={**os.environ,'E2E_FIXTURE_PATH':str(fixture)},capture_output=True,text=True,timeout=240)
        assert result.returncode==0,result.stdout+'\n'+result.stderr
        from pypdf import PdfReader
        reader=PdfReader(print_pdf)
        assert 1 <= len(reader.pages) <= 3
        printed=''.join(page.extract_text() for page in reader.pages)
        for text in ('Browser Debitur','PT Reborn Browser QA','Jl. Kantor Browser QA 12','office@browser-qa.example.com','021-987654'):assert text in printed
    finally:fixture.unlink(missing_ok=True)

def test_chat_attachment_location_idempotency_and_historical_access(live):
    import json
    from pypdf import PdfWriter
    account,_=unit(live);letter=assign(live,account);task=letter['assignment_id'];call=live['call'];db=live['db']
    for role in ('admin', 'officer'):
        inbox = call(role, 'GET', '/chat').json()['items']
        assert any(item['assignment_id'] == task and 'mulai chat' in item['last_message'] for item in inbox)
    photo=io.BytesIO();Image.new('RGB',(32,32),'orange').save(photo,'PNG')
    pdf=io.BytesIO();writer=PdfWriter();writer.add_blank_page(width=100,height=100);writer.write(pdf)
    data={'client_message_id':str(uuid.uuid4()),'text':'Koordinasi kunjungan','location':json.dumps({'latitude':-6.2,'longitude':106.8,'accuracy':10})}
    files=[('files',('photo.png',photo.getvalue(),'image/png')),('files',('document.pdf',pdf.getvalue(),'application/pdf'))]
    message=call('officer','POST',f'/chat/{task}/messages',data=data,files=files).json()
    assert message['sequence']==1 and len(message['attachments'])==2 and 'fingerprint' not in message
    assert all('storage_path' not in item for item in message['attachments'])
    assert call('officer','POST',f'/chat/{task}/messages',data=data,files=files).json()['id']==message['id']
    call('officer','POST',f'/chat/{task}/messages',409,data={**data,'text':'Changed'},files=files)
    assert db.chat_messages.count_documents({'assignment_id':task})==1
    for attachment in message['attachments']:
        path=attachment['url'].removeprefix('/api')
        response=call('admin','GET',path)
        assert response.headers['x-content-type-options']=='nosniff'
        assert response.headers['content-disposition'].startswith('inline' if attachment['kind']=='image' else 'attachment')
        call('foreign','GET',path,404);call('other','GET',path,404)
    assert call('admin','GET',f'/chat/{task}').json()['unread']==1
    assert call('admin','GET','/chat/unread').json()['count']>=1
    call('admin','POST',f'/chat/{task}/read',json={'sequence':9999})
    assert call('admin','GET',f'/chat/{task}').json()['unread']==0
    call('admin','POST',f'/chat/{task}/read',json={'sequence':0})
    assert call('admin','GET',f'/chat/{task}').json()['unread']==0
    call('other','GET',f'/chat/{task}',404);call('foreign','GET',f'/chat/{task}',404)
    inbox=call('officer','GET','/chat').json()['items'];assert any(item['assignment_id']==task for item in inbox)
    call('admin','PATCH',f"/surat-tugas/{letter['id']}/status",json={'status':'selesai'})
    assert call('officer','GET',f'/chat/{task}').json()['assignment']['read_only']
    assert call('officer','POST',f'/chat/{task}/messages',data=data,files=files).json()['id']==message['id']
    call('officer','POST',f'/chat/{task}/messages',409,data={'client_message_id':str(uuid.uuid4()),'text':'Closed'})
    call('officer','GET',message['attachments'][0]['url'].removeprefix('/api'))


def test_chat_concurrent_retry_sequence_and_pagination(live):
    account,_=unit(live);letter=assign(live,account);task=letter['assignment_id'];call=live['call'];db=live['db']
    data={'client_message_id':str(uuid.uuid4()),'text':'Retry in parallel'}
    def send():return requests.post(live['url']+f'/chat/{task}/messages',headers={'Authorization':'Bearer '+live['tokens']['officer']},data=data,timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:send(),range(2)))
    assert [r.status_code for r in responses]==[200,200]
    assert responses[0].json()['id']==responses[1].json()['id']
    assert db.chat_messages.count_documents({'assignment_id':task})==1
    for index in range(4):call('admin','POST',f'/chat/{task}/messages',data={'client_message_id':str(uuid.uuid4()),'text':f'Message {index}'})
    latest=call('officer','GET',f'/chat/{task}?limit=2').json();assert [m['sequence'] for m in latest['messages']]==[4,5] and latest['has_more']
    previous=call('officer','GET',f'/chat/{task}?before=4&limit=2').json();assert [m['sequence'] for m in previous['messages']]==[2,3]
    following=call('officer','GET',f'/chat/{task}?after=1&limit=2').json();assert [m['sequence'] for m in following['messages']]==[2,3] and following['has_more']


def test_chat_audit_failure_rolls_back_message_and_sequence(live):
    account,_=unit(live);letter=assign(live,account);task=letter['assignment_id'];db=live['db'];call=live['call']
    db.block_writes('audit_logs',True)
    try:
        call('admin','POST',f'/chat/{task}/messages',500,data={'client_message_id':str(uuid.uuid4()),'text':'Must roll back'})
        assert db.chat_messages.count_documents({'assignment_id':task})==0
        assert not db.assignments.find_one({'id':task}).get('chat_sequence')
    finally:db.block_writes('audit_logs',False)


def test_chat_invalid_file_and_owner_have_no_upload(live):
    account,_=unit(live);letter=assign(live,account);task=letter['assignment_id'];call=live['call'];db=live['db']
    before=db.upload_intents.count_documents({})
    call('other','POST',f'/chat/{task}/messages',404,data={'client_message_id':str(uuid.uuid4()),'text':'Blocked'})
    call('officer','POST',f'/chat/{task}/messages',400,data={'client_message_id':str(uuid.uuid4())},files={'files':('fake.pdf',b'invalid','application/pdf')})
    call('officer','POST',f'/chat/{task}/messages',400,data={'client_message_id':str(uuid.uuid4())},files={'files':('app.exe',b'fake','application/octet-stream')})
    assert db.upload_intents.count_documents({})==before
    assert db.chat_messages.count_documents({'assignment_id':task})==0

def test_chat_send_racing_close_never_writes_after_closed_state(live):
    account,_=unit(live);letter=assign(live,account);task=letter['assignment_id'];db=live['db']
    def send():return requests.post(live['url']+f'/chat/{task}/messages',headers={'Authorization':'Bearer '+live['tokens']['officer']},data={'client_message_id':str(uuid.uuid4()),'text':'Race with closing'},timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:
        closing=pool.submit(live['call'],'admin','PATCH',f"/surat-tugas/{letter['id']}/status",json={'status':'selesai'})
        pending=pool.submit(send);closing.result();response=pending.result()
    assert response.status_code in {200,409}
    assert db.assignments.find_one({'id':task})['status']=='SELESAI'
    assert db.chat_messages.count_documents({'assignment_id':task})==int(response.status_code==200)
    live['call']('officer','POST',f'/chat/{task}/messages',409,data={'client_message_id':str(uuid.uuid4()),'text':'Definitely closed'})


def test_team_assignment_letter_chat_report_and_access_isolation(live):
    import json
    account, sk = unit(live); call = live['call']; db = live['db']
    ids = [live['users'][role]['id'] for role in ('officer', 'other')]
    letter = call('admin', 'POST', '/penugasan', json={'account_id': account['id'], 'petugas_ids': ids}).json()
    letter=call('owner','POST','/penugasan/'+letter['assignment_id']+'/approve').json()
    task = letter['assignment_id']
    assert letter['petugas_count'] == 2 and letter['petugas_ids'] == ids
    assert db.assignments.count_documents({'account_id': account['id']}) == 1
    assert db.assignment_letters.count_documents({'assignment_id': task}) == 1
    for role in ('officer', 'other'):
        assert any(item['assignment_id'] == task for item in call(role, 'GET', '/my/tugas').json())
        document = call(role, 'GET', f"/my/surat-tugas/{letter['id']}/document").json()
        assert [person['id'] for person in document['data']['officers']] == ids
        assert call(role, 'GET', f'/my/tugas/{account["id"]}').json()['letter_id'] == letter['id']
    for role in ('outsider', 'foreign'):
        call(role, 'GET', f'/chat/{task}', 404)
        # Foreign is an administrator and cannot use the officer endpoint.
        call(role, 'GET', f"/my/surat-tugas/{letter['id']}/document", 403 if role == 'foreign' else 404)
    filtered = call('admin', 'GET', '/surat-tugas?petugas_id='+ids[1]).json()
    assert any(item['id'] == letter['id'] for item in filtered)
    call('admin', 'POST', f'/chat/{task}/messages', data={'client_message_id': str(uuid.uuid4()), 'text': 'Pesan untuk seluruh tim'})
    for role in ('officer', 'other'):
        assert call(role, 'GET', f'/chat/{task}').json()['unread'] == 1
        assert call(role, 'GET', '/chat/unread').json()['count'] >= 1
    call('other', 'POST', f'/chat/{task}/read', json={'sequence': 1})
    assert call('officer', 'GET', f'/chat/{task}').json()['unread'] == 1
    photo = io.BytesIO(); Image.new('RGB', (32,32), 'blue').save(photo, 'PNG')
    message = call('other', 'POST', f'/chat/{task}/messages', data={'client_message_id': str(uuid.uuid4()), 'location': json.dumps({'latitude': -6.2, 'longitude': 106.8})}, files={'files': ('team.png', photo.getvalue(), 'image/png')}).json()
    attachment = message['attachments'][0]['url'][4:]
    call('officer', 'GET', attachment); call('outsider', 'GET', attachment, 404)
    report = call('admin', 'POST', '/laporan', data=report_payload(letter, account), files={'photos': ('proof.png', photo.getvalue(), 'image/png')}).json()
    assert report['petugas_ids'] == ids and report['created_by'] == live['users']['admin']['id']
    call('officer', 'POST', '/laporan', 403, data=report_payload(letter, account))
    assert db.field_reports.count_documents({'assignment_id': task}) == 1
    for role in ('officer', 'other'):
        assert any(item['id'] == report['id'] for item in call(role, 'GET', '/my/riwayat').json())
        call(role, 'GET', report['photos'][0]['url'][4:])
    call('outsider', 'GET', report['photos'][0]['url'][4:], 404)
    call('admin', 'DELETE', '/petugas/'+ids[1], 409)
    call('admin', 'PATCH', f"/surat-tugas/{letter['id']}/status", json={'status': 'selesai'})
    assert call('other', 'GET', f'/chat/{task}').json()['assignment']['read_only']
    verification = call('admin', 'GET', '/public/verify/'+letter['generate_code']).json()
    assert verification['petugas_names'] == ['QA officer', 'QA other']
    assert 'nik' not in verification and 'accounts' not in verification


def test_team_assignment_rejects_foreign_member_without_partial_creation(live):
    account, _ = unit(live)
    live['call']('admin', 'POST', '/penugasan', 404, json={'account_id': account['id'], 'petugas_ids': [live['users']['officer']['id'], live['users']['foreign']['id']]})
    assert live['db'].assignments.count_documents({'account_id': account['id']}) == 0


def test_admin_report_team_payment_pending_late_proof_and_retry(live):
    from pypdf import PdfWriter
    call=live['call'];db=live['db'];account,_=unit(live)
    ids=[live['users'][role]['id'] for role in ('officer','other')]
    letter=call('admin','POST','/penugasan',json={'account_id':account['id'],'petugas_ids':ids}).json()
    letter=call('owner','POST','/penugasan/'+letter['assignment_id']+'/approve').json()
    task=letter['assignment_id'];data=report_payload(letter,account)
    assert any(item['assignment_id']==task for item in call('admin','GET','/laporan-options').json()['items'])
    image=report_files()['photos'][1]
    files=[('photos',('one.png',image,'image/png')),('photos',('two.png',image,'image/png'))]
    call('officer','POST','/laporan',403,data=data,files=files)
    report=call('admin','POST','/laporan',data=data,files=files).json()
    assert report['payment_status']=='MENUNGGU_PEMBAYARAN' and report['payment_amount']=='500000.00'
    assert report['created_role']=='admin' and report['petugas_ids']==ids and len(report['photos'])==2
    assert call('admin','POST','/laporan',data=data,files=files).json()['id']==report['id']
    assert db.field_reports.count_documents({'assignment_id':task})==1
    assert not any(item['assignment_id']==task for item in call('admin','GET','/laporan-options').json()['items'])
    writer=PdfWriter();writer.add_blank_page(width=100,height=100);proof=io.BytesIO();writer.write(proof)
    payment={'payment_request_id':str(uuid.uuid4())};upload={'file':('payment.pdf',proof.getvalue(),'application/pdf')}
    call('officer','POST',f"/laporan/{report['id']}/payment",403,data=payment,files=upload)
    call('foreign','POST',f"/laporan/{report['id']}/payment",404,data=payment,files=upload)
    paid=call('admin','POST',f"/laporan/{report['id']}/payment",data=payment,files=upload).json()
    assert paid['id']==report['id'] and paid['payment_status']=='SELESAI' and paid['payment_amount']=='500000.00'
    assert call('admin','POST',f"/laporan/{report['id']}/payment",data=payment,files=upload).json()['payment_proof']['url']==paid['payment_proof']['url']
    call('admin','POST',f"/laporan/{report['id']}/payment",409,data={'payment_request_id':str(uuid.uuid4())},files=upload)
    assert len(call('admin','GET',f"/laporan/{report['id']}").json()['photos'])==2
    for role in ('officer','other'):
        history=call(role,'GET','/my/riwayat?payment_status=SELESAI').json()
        assert any(item['id']==report['id'] and item['payment_amount']=='500000.00' for item in history)
        assert call(role,'GET',paid['payment_proof']['url'][4:]).content==proof.getvalue()
    call('outsider','GET',paid['payment_proof']['url'][4:],404)
    call('foreign','GET',paid['payment_proof']['url'][4:],404)
    call('admin','PATCH',f"/laporan/{report['id']}/review",json={'catatan_admin':'Tetap selesai'})
    assert db.field_reports.find_one({'id':report['id']})['report_status']=='SELESAI'


def test_admin_report_immediate_payment_and_payment_audit_rollback(live):
    call=live['call'];db=live['db'];account,_=unit(live);letter=assign(live,account)
    photo=report_files()['photos'][1]
    files=[('photos',('evidence.png',photo,'image/png')),('payment_proof',('paid.png',photo,'image/png'))]
    paid=call('admin','POST','/laporan',data=report_payload(letter,account),files=files).json()
    assert paid['payment_status']=='SELESAI' and len(paid['photos'])==1 and paid['payment_proof']
    other,_=unit(live);other_letter=assign(live,other)
    pending=call('admin','POST','/laporan',data=report_payload(other_letter,other),files=report_files()).json()
    db.block_writes('audit_logs',True)
    try:
        call('admin','POST',f"/laporan/{pending['id']}/payment",500,data={'payment_request_id':str(uuid.uuid4())},files={'file':('payment.png',photo,'image/png')})
        stored=db.field_reports.find_one({'id':pending['id']})
        assert stored['payment_status']=='MENUNGGU_PEMBAYARAN' and not stored.get('payment_proof')
        assert db.report_photos.count_documents({'report_id':pending['id'],'kind':'payment'})==0
    finally:db.block_writes('audit_logs',False)


def test_concurrent_team_payment_keeps_one_authoritative_proof(live):
    account,_=unit(live);letter=assign(live,account);call=live['call']
    report=call('admin','POST','/laporan',data=report_payload(letter,account),files=report_files()).json()
    photo=report_files()['photos'][1]
    def pay():return requests.post(live['url']+f"/laporan/{report['id']}/payment",headers={'Authorization':'Bearer '+live['tokens']['admin']},data={'payment_request_id':str(uuid.uuid4())},files={'file':('payment.png',photo,'image/png')},timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:pay(),range(2)))
    assert sorted(response.status_code for response in responses)==[200,409]
    assert live['db'].report_photos.count_documents({'report_id':report['id'],'kind':'payment'})==1


def test_leasing_payment_concurrency_and_private_proof(live):
    call=live['call'];db=live['db']
    leasing=call('admin','POST','/clients',json={'nama_perusahaan':'Leasing pembayaran QA'}).json()
    payload={'client_id':leasing['id'],'payment_amount':'1500000.25','payment_request_id':str(uuid.uuid4())}
    files={'file':('payment.png',report_files()['photos'][1],'image/png')}
    def send():
        return requests.post(live['url']+'/leasing-payments',headers={'Authorization':'Bearer '+live['tokens']['admin']},data=payload,files=files,timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:send(),range(2)))
    assert [r.status_code for r in responses]==[200,200], [r.text for r in responses]
    first=responses[0].json()
    assert first['id']==responses[1].json()['id'] and first['created_at']
    assert first['payment_amount']=='1500000.25'
    assert db.documents.count_documents({'kind':'leasing_payment','client_id':leasing['id']})==1
    listed=call('admin','GET','/leasing-payments').json()['items']
    assert any(item['id']==first['id'] for item in listed)
    assert not call('foreign','GET','/leasing-payments').json()['items']
    assert call('admin','GET',first['url'][4:]).content==files['file'][1]
    for role in ('officer','foreign'):
        call(role,'GET',first['url'][4:],404)
    call('officer','GET','/leasing-payments',403)
    call('foreign','POST','/leasing-payments',404,data=payload,files=files)
    call('admin','POST','/leasing-payments',409,data={**payload,'payment_amount':'2'},files=files)
    call('admin','DELETE',f"/clients/{leasing['id']}",409)


def test_leasing_payment_audit_failure_rolls_back_record(live):
    call=live['call'];db=live['db']
    leasing=call('admin','POST','/clients',json={'nama_perusahaan':'Leasing rollback QA'}).json()
    db.block_writes('audit_logs',True)
    try:
        call('admin','POST','/leasing-payments',500,data={'client_id':leasing['id'],'payment_amount':'100','payment_request_id':str(uuid.uuid4())},files={'file':('paid.png',report_files()['photos'][1],'image/png')})
        assert db.documents.count_documents({'kind':'leasing_payment','client_id':leasing['id']})==0
    finally:
        db.block_writes('audit_logs',False)


def test_finance_union_totals_dates_pagination_and_exports(live):
    from decimal import Decimal
    from openpyxl import load_workbook
    from pypdf import PdfReader
    db=live['db'];call=live['call'];company=live['users']['admin']['company_id']
    def income(nominal,stamp,tenant=company,kind='leasing_payment'):
        db.documents.insert_one({'id':uuid.uuid4().hex,'company_id':tenant,'kind':kind,'leasing_name':'Leasing Finance QA','payment_amount':nominal,'created_at':stamp,'url':'/api/files/in.png'})
    def expense(nominal,stamp,**extra):
        db.field_reports.insert_one({'id':uuid.uuid4().hex,'company_id':company,'payment_status':'SELESAI','payment_amount':nominal,'payment_recorded_at':stamp,'created_at':'2020-01-01T00:00:00+00:00','team_names':['Satu','Dua'],'payment_proof':{'url':'/api/files/out.png'},**extra})
    income('1000000.25','2035-01-01T17:00:00+00:00')
    income('500000.10','2035-01-02T16:59:59+00:00')
    expense('250000.15','2035-01-02T03:00:00+00:00')
    income('999','2035-01-02T17:00:00+00:00') # Following WIB day, excluded.
    income('999','2035-01-01T16:59:59+00:00') # Previous WIB day, excluded.
    income('999','2035-01-02T10:00:00+00:00',tenant='foreign-company')
    income('999','2035-01-02T10:00:00+00:00',kind='surat_kuasa')
    expense('999','2035-01-02T10:00:00+00:00',payment_status='MENUNGGU_PEMBAYARAN')
    expense('999','2035-01-02T10:00:00+00:00',payment_proof=None)
    query='?start=2035-01-02&end=2035-01-02'
    report=call('admin','GET','/finance'+query+'&limit=1').json()
    assert report['total']==3 and len(report['items'])==1 and report['next_page']==2
    assert report['summary']=={'income':'1500000.35','expense':'250000.15','balance':'1250000.20'}
    page2=call('admin','GET','/finance'+query+'&limit=1&page=2').json()
    assert page2['items'][0]['id']!=report['items'][0]['id'] and page2['summary']==report['summary']
    outgoing=call('admin','GET','/finance'+query+'&kind=PENGELUARAN').json()
    assert outgoing['total']==1 and outgoing['items'][0]['party']=='Satu, Dua'
    assert outgoing['summary']['balance']=='-250000.15'
    for role in ('officer','outsider'):
        call(role,'GET','/finance'+query,403)
        call(role,'GET','/finance/export/xlsx'+query,403)
    foreign=call('foreign','GET','/finance'+query).json()
    assert foreign['total']==1 and foreign['summary']['income']=='999.00'
    sheet=load_workbook(io.BytesIO(call('admin','GET','/finance/export/xlsx'+query).content)).active
    assert sheet.max_row==12 and Decimal(str(sheet['B6'].value))==Decimal('1250000.20')
    document=PdfReader(io.BytesIO(call('admin','GET','/finance/export/pdf'+query).content))
    assert '1.250.000,20' in document.pages[0].extract_text()


def test_pks_dates_history_private_files_and_concurrent_retry(live):
    from pypdf import PdfWriter
    call=live['call'];db=live['db'];officer=live['users']['officer'];company=officer['company_id']
    writer=PdfWriter();writer.add_blank_page(width=595,height=842);output=io.BytesIO();writer.write(output)
    files={'file':('PKS.pdf',output.getvalue(),'application/pdf')}
    before=call('admin','GET','/pks/officers?search=QA%20officer').json()['items']
    assert len(before)==1 and before[0]['pks_status']=='BELUM_MEMILIKI'
    preview=call('admin','GET','/pks/period?duration=3&unit=TAHUN').json()
    payload={'officer_id':officer['id'],'duration':'3','unit':'TAHUN','submission_id':str(uuid.uuid4()),'start_date':'2000-01-01','end_date':'2099-12-31'}
    def send():return requests.post(live['url']+'/pks',headers={'Authorization':'Bearer '+live['tokens']['admin']},data=payload,files=files,timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:send(),range(2)))
    assert [r.status_code for r in responses]==[200,200],[r.text for r in responses]
    record=responses[0].json()
    assert record['id']==responses[1].json()['id']
    assert record['start_date']==preview['start_date'] and record['end_date']==preview['end_date']
    assert record['status']=='AKTIF' and db.documents.count_documents({'kind':'pks','officer_id':officer['id']})==1
    assert call('admin','GET',record['url'][4:]).content==output.getvalue()
    for role in ('officer','other','foreign'):
        call(role,'GET',record['url'][4:],404)
    call('officer','GET','/pks',403);call('officer','POST','/pks',403,data=payload,files=files)
    call('foreign','POST','/pks',404,data=payload,files=files)
    assert not call('foreign','GET','/pks').json()['items']
    call('admin','POST','/pks',409,data={**payload,'duration':'2'},files=files)
    renewed=call('admin','POST','/pks',data={**payload,'duration':'4','submission_id':str(uuid.uuid4())},files=files).json()
    assert renewed['id']!=record['id'] and db.documents.count_documents({'kind':'pks','officer_id':officer['id']})==2
    assert call('admin','GET',record['url'][4:]).content==output.getvalue()
    after=call('admin','GET','/pks/officers?search=QA%20officer').json()['items'][0]
    assert after['pks_status']=='AKTIF' and after['pks_end_date']==renewed['end_date']
    day=datetime.now(timezone(timedelta(hours=7))).date()
    for days,label in [(30,'MENDEKATI_EXPIRED'),(0,'MENDEKATI_EXPIRED'),(-1,'EXPIRED')]:
        identity=uuid.uuid4().hex
        db.documents.insert_one({'id':identity,'company_id':company,'kind':'pks','officer_id':live['users']['other']['id'],'officer_name':'Boundary PKS','start_date':(day-timedelta(days=365)).isoformat(),'end_date':(day+timedelta(days=days)).isoformat(),'created_at':datetime.now(timezone.utc).isoformat()})
        listed=call('admin','GET','/pks?state='+label).json()['items']
        assert any(row['id']==identity and row['status']==label for row in listed)
    db.block_writes('audit_logs',True)
    try:
        call('admin','POST','/pks',500,data={**payload,'officer_id':live['users']['outsider']['id'],'submission_id':str(uuid.uuid4())},files=files)
        assert db.documents.count_documents({'kind':'pks','officer_id':live['users']['outsider']['id']})==0
    finally:db.block_writes('audit_logs',False)


def test_owner_approval_blocks_execution_and_stamp_until_approved(live):
    call=live['call'];db=live['db'];account,_=unit(live)
    letter=call('admin','POST','/penugasan',json={'account_id':account['id'],'petugas_id':live['users']['officer']['id']}).json()
    task=letter['assignment_id'];lid=letter['id']
    assert letter['status']=='menunggu_approval' and letter['document_snapshot']['company']['director_signature']==''
    assert db.assignments.find_one({'id':task})['status']=='MENUNGGU_APPROVAL'
    count=call('owner','GET','/approvals/pending-count').json()['count']
    assert count>=1
    draft=call('admin','GET',f'/surat-tugas/{lid}/document').json()
    assert not draft['director_signature'] and not draft['is_finalized'] and not draft['data']['company'].get('director_signature')
    call('admin','POST',f'/penugasan/{task}/approve',403)
    call('foreign','POST',f'/penugasan/{task}/approve',403)
    call('admin','POST',f'/surat-tugas/{lid}/finalize',403)
    call('admin','PATCH',f'/surat-tugas/{lid}/status',403,json={'status':'aktif'})
    call('admin','PATCH',f'/surat-tugas/{lid}/status',403,json={'status':'selesai'})
    call('officer','GET',f'/my/tugas/{account["id"]}',404)
    call('officer','GET',f'/my/surat-tugas/{lid}/document',403)
    call('officer','GET',f'/chat/{task}',404)
    call('admin','POST','/laporan',409,data=report_payload(letter,account),files=report_files())
    assert call('admin','GET',f'/public/verify/{letter["generate_code"]}').json()['valid'] is False
    db.block_writes('audit_logs',True)
    try:
        call('owner','POST',f'/penugasan/{task}/approve',500)
        assert db.assignments.find_one({'id':task})['status']=='MENUNGGU_APPROVAL'
        assert not call('admin','GET',f'/surat-tugas/{lid}/document').json()['director_signature']
    finally:db.block_writes('audit_logs',False)
    def approve():return requests.post(live['url']+f'/penugasan/{task}/approve',headers={'Authorization':'Bearer '+live['tokens']['owner']},timeout=30)
    with ThreadPoolExecutor(max_workers=2) as pool:responses=list(pool.map(lambda _:approve(),range(2)))
    assert [response.status_code for response in responses]==[200,200],[r.text for r in responses]
    assert call('owner','GET','/approvals/pending-count').json()['count']==count-1
    signed=call('owner','GET',f'/surat-tugas/{lid}/document').json()
    assert signed['is_finalized'] and signed['director_signature'].startswith('data:image/png;base64,')
    call('officer','GET',f'/my/tugas/{account["id"]}')
    call('officer','GET',f'/my/surat-tugas/{lid}/document')
    assert call('admin','GET',f'/public/verify/{letter["generate_code"]}').json()['valid'] is True
    db.companies.update_one({'id':account['company_id']},{'$set':{'director_signature':'changed-after-approval'}})
    assert call('admin','GET',f'/surat-tugas/{lid}/document').json()['director_signature']==signed['director_signature']
    db.companies.update_one({'id':account['company_id']},{'$set':{'director_signature':signed['director_signature']}})
    assert db.audit_logs.count_documents({'entity':'assignment','entity_id':task,'activity':'Menyetujui surat tugas'})==1


def test_owner_admin_creation_and_monthly_cash_cards(live):
    from decimal import Decimal
    call=live['call'];owner=live['users']['owner']
    password=secrets.token_urlsafe(18);email=uuid.uuid4().hex+'@example.com'
    body={'name':'QA Admin from owner','email':email,'password':password}
    call('admin','POST','/admin-users',403,json=body)
    created=call('owner','POST','/admin-users',json=body).json()
    assert created['role']=='admin' and created['company_id']==owner['company_id'] and 'password_hash' not in created
    call('owner','POST','/admin-users',422,json={**body,'email':uuid.uuid4().hex+'@example.com','role':'owner'})
    call('admin','GET','/admin-users',403)
    assert any(row['id']==created['id'] for row in call('owner','GET','/admin-users').json()['items'])
    before=call('owner','GET','/dashboard/finance').json()
    foreign=call('foreign','GET','/dashboard/finance').json()
    leasing=call('owner','POST','/clients',json={'nama_perusahaan':'Monthly Leasing QA'}).json()
    call('owner','POST','/leasing-payments',data={'client_id':leasing['id'],'payment_amount':'1200000.25','payment_request_id':str(uuid.uuid4())},files={'file':('income.png',report_files()['photos'][1],'image/png')})
    account,_=unit(live);letter=assign(live,account)
    report=call('admin','POST','/laporan',data=report_payload(letter,account),files=report_files()).json()
    waiting=call('admin','GET','/dashboard/finance').json()
    assert waiting['expense']==before['expense']
    call('admin','POST',f'/laporan/{report["id"]}/payment',data={'payment_request_id':str(uuid.uuid4())},files={'file':('out.png',report_files()['photos'][1],'image/png')})
    after=call('owner','GET','/dashboard/finance').json()
    assert Decimal(after['income'])-Decimal(before['income'])==Decimal('1200000.25')
    assert Decimal(after['expense'])-Decimal(before['expense'])==Decimal('500000.00')
    assert Decimal(after['profit'])-Decimal(before['profit'])==Decimal('700000.25')
    assert call('foreign','GET','/dashboard/finance').json()==foreign


def test_owner_cli_recovery_and_real_login(live):
    email='owner-recovery-'+uuid.uuid4().hex+'@example.com'
    original=secrets.token_urlsafe(24);replacement=secrets.token_urlsafe(24)
    command=[sys.executable,str(live['root']/'backend/manage.py'),'bootstrap-owner','--name','Recovery QA','--email',email,'--admin-email',live['users']['admin']['email']]
    def provision(password,reset=False):
        result=subprocess.run(command+(['--reset-password'] if reset else []),cwd=live['root'],env={**live['env'],'BOOTSTRAP_OWNER_PASSWORD':password},capture_output=True,text=True,timeout=90)
        assert result.returncode==0,result.stdout+result.stderr
        assert password not in result.stdout+result.stderr
    provision(original)
    first=requests.post(live['url']+'/auth/login',json={'email':email,'password':original},timeout=30)
    assert first.status_code==200 and first.json()['user']['role']=='owner'
    token=list(first.cookies.values())[0]
    provision(replacement,True)
    assert requests.get(live['url']+'/auth/me',headers={'Authorization':'Bearer '+token},timeout=30).status_code==401
    assert requests.post(live['url']+'/auth/login',json={'email':email,'password':original},timeout=30).status_code==401
    after=requests.post(live['url']+'/auth/login',json={'email':email,'password':replacement},timeout=30)
    assert after.status_code==200 and after.json()['user']['role']=='owner'

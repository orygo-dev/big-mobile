from decimal import Decimal
from io import BytesIO
from unittest.mock import AsyncMock
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from test_local_regressions import local_api, server, ADMIN, OFFICER, bearer
from finance import date_bounds, normalize
from finance_exports import render_pdf, render_xlsx


def result():
    return {'count':3, 'totals':{'PEMASUKAN':Decimal('1000000.25'),'PENGELUARAN':Decimal('250000.10')},'rows':[
        ('income','2026-10-17T03:00:00+00:00','PEMASUKAN',Decimal('1000000.25'),{'leasing_name':'=1+1','url':'/api/files/in.png'}),
        ('expense','2026-10-17T02:00:00+00:00','PENGELUARAN',Decimal('250000.10'),{'team_names':['Petugas Satu','Petugas Dua'],'catatan':'=HYPERLINK("bad")','payment_proof':{'url':'/api/files/out.png'}})]}


def test_report_scope_date_bounds_and_full_totals(local_api):
    api, db = local_api; db.finance_report=AsyncMock(return_value=result())
    response=api.get('/api/finance?start=2026-10-17&end=2026-10-17&limit=1')
    assert response.status_code==200
    body=response.json()
    assert body['summary']=={'income':'1000000.25','expense':'250000.10','balance':'750000.15'}
    assert body['total']==3 and body['next_page']==2
    assert body['items'][1]['party']=='Petugas Satu, Petugas Dua'
    assert body['items'][0]['id']!=body['items'][1]['id']
    db.finance_report.assert_awaited_once_with(ADMIN['company_id'],'2026-10-16T17:00:00+00:00','2026-10-17T17:00:00+00:00','',1,1,False)


@pytest.mark.parametrize('query',['start=bad','start=2026-10-18&end=2026-10-17','end=2026-02-30','kind=OTHER','page=0','start=20261017'])
def test_invalid_filters_rejected(local_api,query):
    api, db=local_api;db.finance_report=AsyncMock()
    assert api.get('/api/finance?'+query).status_code in (400,422)
    db.finance_report.assert_not_called()


def test_officer_cannot_read_or_export_finance(local_api):
    api, db=local_api
    server.app.dependency_overrides.pop(server.admin_required);db.users.find_one.return_value=OFFICER
    for path in ('/finance','/finance/export/pdf','/finance/export/xlsx'):
        assert api.get('/api'+path,headers=bearer(OFFICER)).status_code==403


def test_export_all_filtered_rows_not_one_page(local_api):
    api, db=local_api;db.finance_report=AsyncMock(return_value=result())
    response=api.get('/api/finance/export/xlsx?kind=PEMASUKAN')
    assert response.status_code==200
    assert response.headers['cache-control']=='private, no-store'
    assert 'attachment;' in response.headers['content-disposition']
    db.finance_report.assert_awaited_once_with(ADMIN['company_id'],None,None,'PEMASUKAN',1,10000,True)
    sheet=load_workbook(BytesIO(response.content)).active
    assert sheet['C10'].value=='=1+1' and sheet['C10'].data_type=='s'
    assert sheet['D11'].data_type=='s' and sheet['E10'].data_type=='n'
    assert sheet['E10'].value==1000000.25
    assert sheet['B6'].value==750000.15
    assert sheet['A10'].value.hour==10 and sheet.freeze_panes=='A10'


def test_export_limit_returns_clear_error(local_api):
    api,db=local_api;db.finance_report=AsyncMock(side_effect=ValueError('EXPORT_LIMIT'))
    response=api.get('/api/finance/export/pdf')
    assert response.status_code==400 and '10.000' in response.json()['detail']


def test_pdf_paginated_headers_escaped_notes_and_totals():
    data=normalize(result());data['items']*=30;data['total']=60
    data['items'][1]['description']='Catatan <aman> & pembayaran '+('uraian '*50)
    document=PdfReader(BytesIO(render_pdf(data,'PT. BERKAT INTIM GENERASI','Semua transaksi | WIB')))
    assert len(document.pages)>1
    text=' '.join(page.extract_text() for page in document.pages)
    assert '750.000,15' in text and 'Catatan <aman>' in text
    assert all('Tanggal (WIB)' in page.extract_text() for page in document.pages[:-1])


def test_empty_exports_remain_valid():
    data=normalize({'count':0,'totals':{},'rows':[]})
    assert len(PdfReader(BytesIO(render_pdf(data,'Perusahaan','Semua transaksi'))).pages)==1
    assert load_workbook(BytesIO(render_xlsx(data,'Perusahaan','Semua transaksi'))).active['B6'].value==0


def test_pdf_keeps_long_notes_across_pages():
    data=normalize(result());data['items']=data['items'][1:];data['total']=1
    data['items'][0]['description']=('catatan panjang '*240)+' AKHIR CATATAN'
    document=PdfReader(BytesIO(render_pdf(data,'Perusahaan','Semua transaksi')))
    assert 'AKHIR CATATAN' in ' '.join(page.extract_text() for page in document.pages)

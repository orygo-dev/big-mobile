"""Server-side PDF and XLSX reports with the same rows and totals as the ledger."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import re
from xml.sax.saxutils import escape
from finance import WIB


def text(value, maximum=4000):
    return ' '.join(str(value or '').split())[:maximum]


def local_date(value):
    return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(WIB).strftime('%d/%m/%Y %H:%M')


def money(value):
    return 'Rp '+f'{Decimal(value):,.2f}'.replace(',', '_').replace('.', ',').replace('_', '.')


def render_pdf(data, company, period):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=32, leftMargin=32, topMargin=30, bottomMargin=32, title='Laporan Keuangan', author=text(company))
    normal = ParagraphStyle('FinanceText', fontName='Helvetica', fontSize=8, leading=11, textColor=colors.HexColor('#183153'))
    heading = ParagraphStyle('FinanceHeading', parent=normal, fontName='Helvetica-Bold', fontSize=18, leading=23)
    def p(value):return Paragraph(escape(text(value)), normal)
    story = [Paragraph('LAPORAN KEUANGAN', heading), p(company), p(period), Spacer(1, 12)]
    summary = Table([[p('TOTAL PEMASUKAN'), p('TOTAL PENGELUARAN'), p('SELISIH')], [p(money(data['summary']['income'])), p(money(data['summary']['expense'])), p(money(data['summary']['balance']))]], colWidths=[doc.width/3]*3)
    summary.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#edf5ff')),('BOX',(0,0),(-1,-1),0.5,colors.HexColor('#d4e3f8')),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    story += [summary, Spacer(1, 14)]
    rows = [[p(label) for label in ['Tanggal (WIB)', 'Jenis', 'Leasing / Tim', 'Keterangan', 'Pemasukan (Rp)', 'Pengeluaran (Rp)']]]
    for row in data['items']:
        incoming = row['type']=='PEMASUKAN'
        rows.append([p(local_date(row['date'])), p(row['type'].title()), p(row['party']), p(row['description']), p(money(row['amount']) if incoming else '-'), p(money(row['amount']) if not incoming else '-')])
    if not data['items']:rows.append([p('Tidak ada transaksi pada filter ini.'), '', '', '', '', ''])
    widths = [92,75,145,doc.width-542,115,115]
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign='LEFT', splitInRow=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#daeaff')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),0.8,colors.HexColor('#1071d8')),('LINEBELOW',(0,1),(-1,-1),0.35,colors.HexColor('#dce5ef')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f8fbff')]),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    story += [table, Spacer(1, 10), p(f"Jumlah transaksi: {data['total']}. Pengeluaran hanya mencakup pembayaran petugas yang sudah memiliki bukti pembayaran.")]
    def footer(canvas, document):
        canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#64748b'));canvas.drawString(32,16,'Laporan keuangan | '+datetime.now(WIB).strftime('%d/%m/%Y %H:%M WIB'));canvas.drawRightString(document.pagesize[0]-32,16,f'Halaman {document.page}');canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def render_xlsx(data, company, period):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    output = BytesIO();workbook = Workbook();sheet = workbook.active;sheet.title='Laporan Keuangan'
    def literal(cell, value):
        # User-controlled names/notes must remain text, including strings beginning with '='.
        cell.value=re.sub(r'[\x00-\x08\x0b-\x0c\x0e-\x1f]', '', str(value or ''))[:32767];cell.data_type='s'
    sheet.merge_cells('A1:G1');literal(sheet['A1'], 'LAPORAN KEUANGAN - '+text(company));sheet['A1'].font=Font(size=16,bold=True,color='075AB8');sheet.row_dimensions[1].height=26
    sheet.merge_cells('A2:G2');literal(sheet['A2'],period);sheet.row_dimensions[2].height=25
    for row,label,key in [(4,'Total Pemasukan','income'),(5,'Total Pengeluaran','expense'),(6,'Selisih','balance')]:
        sheet.cell(row,1,label);sheet.cell(row,2,Decimal(data['summary'][key]));sheet.cell(row,2).number_format='"Rp" #,##0.00;[Red]("Rp" #,##0.00)';sheet.cell(row,1).font=Font(bold=True);sheet.cell(row,2).fill=PatternFill('solid',fgColor='EAF3FF')
    sheet.merge_cells('A7:G7');sheet['A7']=f"{data['total']} transaksi | Waktu WIB | Pengeluaran berdasarkan pembayaran petugas yang telah selesai"
    headers=['Tanggal (WIB)','Jenis','Leasing / Tim','Keterangan','Pemasukan (Rp)','Pengeluaran (Rp)','Referensi Sumber']
    for column,label in enumerate(headers,1):
        cell=sheet.cell(9,column,label);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='0967D2')
    for number,row in enumerate(data['items'],10):
        sheet.cell(number,1,datetime.fromisoformat(row['date'].replace('Z','+00:00')).astimezone(WIB).replace(tzinfo=None)).number_format='dd/mm/yyyy hh:mm'
        literal(sheet.cell(number,2),row['type']);literal(sheet.cell(number,3),row['party']);literal(sheet.cell(number,4),row['description']);literal(sheet.cell(number,7),row['id'])
        sheet.cell(number,5,Decimal(row['amount']) if row['type']=='PEMASUKAN' else 0);sheet.cell(number,6,Decimal(row['amount']) if row['type']=='PENGELUARAN' else 0)
        for column in (5,6):sheet.cell(number,column).number_format='#,##0.00;[Red](#,##0.00)'
        for cell in sheet[number]:
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if number%2==0:cell.fill=PatternFill('solid',fgColor='F0F6FF')
    for column,width in {'A':23,'B':19,'C':35,'D':60,'E':23,'F':23,'G':55}.items():sheet.column_dimensions[column].width=width
    sheet.freeze_panes='A10';sheet.auto_filter.ref=f'A9:G{max(9,len(data["items"])+9)}';sheet.sheet_view.showGridLines=False
    sheet.print_title_rows='1:9';sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    workbook.save(output)
    return output.getvalue()
